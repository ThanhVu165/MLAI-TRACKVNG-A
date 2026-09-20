import io
import unicodedata

from docx import Document

from corpus.extract_doc import extract_document, normalize_pages


def test_normalize_removes_repeated_marginals_and_preserves_legal_markers() -> None:
    pages = [
        "TRƯỜNG ĐẠI HỌC\nĐiều 1. Phạm vi\nNội dung bị ngắt\ngiữa câu\nTrang",
        "TRƯỜNG ĐẠI HỌC\nĐiều 2. Điều kiện\n1. Khoản thứ nhất\na) Điểm chi tiết\nTrang",
        "TRƯỜNG ĐẠI HỌC\nĐiều 3. Hiệu lực\nVăn bản có hiệu lực.\nTrang",
    ]

    text = normalize_pages(pages)

    assert "TRƯỜNG ĐẠI HỌC" not in text and "Trang" not in text
    assert text.splitlines()[0] == "Điều 1. Phạm vi"
    assert "Nội dung bị ngắt giữa câu" in text
    assert "\n1. Khoản thứ nhất\na) Điểm chi tiết" in text


def test_extract_docx_keeps_numbering_and_normalizes_nfc() -> None:
    document = Document()
    document.add_paragraph("Điều 8. Rút học phần")
    document.add_paragraph("1. Sinh viên nộp đơn trực tuyến.")
    document.add_paragraph("a) Thời hạn áp dụng.")
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_document(buffer.getvalue(), "quy-dinh.docx")

    assert text.splitlines() == [
        "Điều 8. Rút học phần",
        "1. Sinh viên nộp đơn trực tuyến.",
        "a) Thời hạn áp dụng.",
    ]
    assert text == unicodedata.normalize("NFC", text)
